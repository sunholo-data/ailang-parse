# /// script
# requires-python = ">=3.10"
# dependencies = ["openpyxl>=3.1"]
# ///
"""Build the XLSX colour fixtures with KNOWN expected colours, and check them.

Colour extraction (design_docs/planned/v0_43_0/v0_43_0_colour_extraction.md)
reached a release-candidate state with no fixture that contained a single cell
fill: every xlsx in data/test_files had zero solid fills, so the office suite
scored 100% on a parser that dropped every colour. These files exist so that
cannot happen again, and their expected values are computed HERE, in Python,
independently of the AILANG implementation — a golden generated from our own
output would only prove the parser agrees with itself.

Theme tints are recomputed with colorsys from the theme openpyxl wrote, not
copied from the parser.

Usage:
    uv run benchmarks/create_colour_fixtures.py            # (re)build fixtures
    uv run benchmarks/create_colour_fixtures.py --verify   # parse + compare
"""

from __future__ import annotations

import colorsys
import json
import os
import re
import subprocess
import sys
import zipfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
from openpyxl.styles import Color, Font, GradientFill, PatternFill

REPO = Path(__file__).resolve().parent.parent
OUT = REPO / "data" / "test_files"


def solid(rgb: str) -> PatternFill:
    return PatternFill(fill_type="solid", fgColor=rgb)


# --- independent oracle ------------------------------------------------------

def theme_colours(xlsx: Path) -> dict[str, str]:
    """Scheme colours by element name, read straight from theme1.xml."""
    xml = zipfile.ZipFile(xlsx).read("xl/theme/theme1.xml").decode()
    scheme = re.search(r"<a:clrScheme.*?</a:clrScheme>", xml, re.S).group(0)
    out = {}
    for name, body in re.findall(r"<a:(\w+)>(.*?)</a:\1>", scheme, re.S):
        m = re.search(r'srgbClr val="([0-9A-Fa-f]{6})"', body) or \
            re.search(r'lastClr="([0-9A-Fa-f]{6})"', body)
        if m:
            out[name] = "#" + m.group(1).lower()
    return out


SLOTS = ["lt1", "dk1", "lt2", "dk2", "accent1", "accent2", "accent3",
         "accent4", "accent5", "accent6", "hlink", "folHlink"]


def tint(hexcol: str, t: float) -> str:
    r, g, b = (int(hexcol[i:i + 2], 16) / 255 for i in (1, 3, 5))
    h, l, s = colorsys.rgb_to_hls(r, g, b)
    l = l * (1 + t) if t < 0 else l * (1 - t) + t
    r, g, b = colorsys.hls_to_rgb(h, l, s)
    return "#" + "".join(f"{int(round(c * 255)):02x}" for c in (r, g, b))


# --- fixtures ----------------------------------------------------------------
#
# Each expectation is (sheet, row, col, text, fill, color), 1-based row/col as
# Excel shows them. Row 1 is the table header in docparse output.

def build_fills(path: Path) -> list:
    wb = Workbook()
    ws = wb.active
    ws.title = "Fills"
    ws.append(["Case", "Value"])
    rows = [
        ("rgb solid", "yes", solid("FFFF00"), None),
        ("theme accent1 +40%", "t", PatternFill("solid", fgColor=Color(theme=4, tint=0.3999755851924192)), None),
        ("theme accent1 -25%", "t", PatternFill("solid", fgColor=Color(theme=4, tint=-0.249977111117893)), None),
        # 2/3 are lt2/dk2: the slots whose index order is swapped against the
        # element order in theme1.xml. A positional lookup gets these wrong.
        ("theme 2 (lt2)", "t", PatternFill("solid", fgColor=Color(theme=2)), None),
        ("theme 3 (dk2)", "t", PatternFill("solid", fgColor=Color(theme=3)), None),
        ("indexed 22", "i", PatternFill("solid", fgColor=Color(indexed=22)), None),
        ("font colour", "red text", None, Font(color="FF0000")),
        ("font theme accent2", "t", None, Font(color=Color(theme=5))),
        # Explicitly black == the Normal font's colour: not a signal, so "".
        ("font explicit default", "t", None, Font(color=Color(theme=1))),
        ("pattern darkGray", "p", PatternFill("darkGray", fgColor="00B050"), None),
        ("gradient", "g", GradientFill(stop=("FF0000", "0000FF")), None),
        ("empty styled", None, solid("FFC000"), None),
        ("no fill", "plain", None, None),
    ]
    for label, value, fill, font in rows:
        ws.append([label, value])
        c = ws.cell(row=ws.max_row, column=2)
        if fill is not None:
            c.fill = fill
        if font is not None:
            c.font = font
    wb.save(path)

    th = theme_colours(path)
    slot = lambda i: th[SLOTS[i]]
    e = lambda r, text, fill="", color="": ("Fills", r, 2, text, fill, color)
    return [
        e(2, "yes", "#ffff00"),
        e(3, "t", tint(slot(4), 0.3999755851924192)),
        e(4, "t", tint(slot(4), -0.249977111117893)),
        e(5, "t", slot(2)),
        e(6, "t", slot(3)),
        e(7, "i", "#c0c0c0"),
        e(8, "red text", "", "#ff0000"),
        e(9, "t", "", slot(5)),
        e(10, "t", "", ""),
        e(11, "p", "#00b050"),
        e(12, "g", ""),
        e(13, "", "#ffc000"),
        e(14, "plain"),
    ]


def build_form(path: Path) -> list:
    """A realistic form: the EMPTY yellow cells are the fields to fill."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Supplier form"
    ws["A1"] = "Supplier onboarding"
    ws.merge_cells("A1:B1")
    ws["A1"].fill = solid("1F4E78")
    ws["A1"].font = Font(color="FFFFFF", bold=True)
    labels = ["Company name", "VAT number", "Contact email", "Start date", "Approved?"]
    for i, label in enumerate(labels, start=2):
        ws.cell(row=i, column=1, value=label)
        ws.cell(row=i, column=2).fill = solid("FFFF00")   # input cell, no value
    ws.cell(row=7, column=1, value="Office use only")
    ws.cell(row=7, column=2, value="N/A").fill = solid("D9D9D9")
    wb.save(path)

    exp = [("Supplier form", 1, 1, "Supplier onboarding", "#1f4e78", "#ffffff")]
    exp += [("Supplier form", i, 2, "", "#ffff00", "") for i in range(2, 7)]
    exp += [("Supplier form", 7, 2, "N/A", "#d9d9d9", "")]
    exp += [("Supplier form", i, 1, l, "", "") for i, l in enumerate(labels, start=2)]
    return exp


def build_condfmt(path: Path) -> list:
    """Colour that only a viewer computes: must WARN, not be invented."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Status"
    ws.append(["Item", "Amount", "Score"])
    for i, (a, s) in enumerate([(50, 1), (150, 5), (90, 3), (300, 9), (10, 2)]):
        ws.append([f"Item {i + 1}", a, s])
    ws.conditional_formatting.add("B2:B6", CellIsRule(operator="greaterThan", formula=["100"], fill=solid("FFC7CE")))
    ws.conditional_formatting.add("C2:C6", ColorScaleRule(start_type="min", start_color="F8696B", end_type="max", end_color="63BE7B"))
    wb.save(path)
    # No direct colour anywhere.
    return [("Status", r, c, None, "", "") for r in range(1, 7) for c in (1, 2, 3)]


def build_no_theme(path: Path) -> list:
    """build_fills with xl/theme/theme1.xml removed: theme references cannot
    be resolved, so they must come back "" AND say so, not guess a colour."""
    tmp = path.with_suffix(".tmp.xlsx")
    exp = build_fills(tmp)
    with zipfile.ZipFile(tmp) as src, zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            if item.filename != "xl/theme/theme1.xml":
                dst.writestr(item, src.read(item.filename))
    tmp.unlink()
    themed = {3, 4, 5, 6, 9}   # rows whose colour is a theme reference
    return [(s_, r, c, t, "" if r in themed else f, "" if r in themed else col)
            for (s_, r, c, t, f, col) in exp]


FIXTURES = {
    "xlsx_fills.xlsx": build_fills,
    "xlsx_no_theme.xlsx": build_no_theme,
    "xlsx_form_template.xlsx": build_form,
    "xlsx_condfmt.xlsx": build_condfmt,
}

EXPECTED_WARNINGS = {
    "xlsx_fills.xlsx": [r"pattern", r"gradient"],
    # 6 = the four theme fills + the two theme fonts; NOT every styled cell
    # (the default font's theme reference is not author colour).
    "xlsx_no_theme.xlsx": [r"6 cell\(s\) use theme colours", r"pattern", r"gradient"],
    "xlsx_form_template.xlsx": [],
    "xlsx_condfmt.xlsx": [r"conditional formatting"],
}


# --- verification ------------------------------------------------------------

def cell_grid(doc: dict) -> dict[str, list[list[dict]]]:
    """sheet name -> rows (header first) of {text, fill, color}."""
    out = {}
    def norm(c):
        if isinstance(c, str):
            return {"text": c, "fill": "", "color": ""}
        return {"text": c.get("text", ""), "fill": c.get("fill", ""), "color": c.get("color", "")}
    for b in doc["document"]["blocks"]:
        if b.get("type") == "section" and b.get("kind") == "sheet":
            for t in b["blocks"]:
                if t.get("type") == "table":
                    out[b["name"]] = [[norm(c) for c in t["headers"]]] + \
                                     [[norm(c) for c in r] for r in t["rows"]]
    return out


def verify() -> int:
    outdir = REPO / "docparse" / "data"
    files = [OUT / f for f in FIXTURES]
    subprocess.run(
        ["ailang", "run", "--entry", "main", "--caps", "IO,FS,Env",
         "--max-recursion-depth", "50000", "--batch", "docparse/main.ail"]
        + [str(f) for f in files],
        cwd=REPO, check=True, capture_output=True,
        env={**os.environ, "DOCPARSE_OUTPUT_DIR": str(outdir)})
    failures = 0
    for name in FIXTURES:
        doc = json.load(open(outdir / f"{name}.json"))
        grid = cell_grid(doc)
        exp = EXPECTATIONS[name]
        for sheet, r, c, text, fill, color in exp:
            try:
                got = grid[sheet][r - 1][c - 1]
            except (KeyError, IndexError):
                print(f"FAIL {name} {sheet}!R{r}C{c}: cell missing")
                failures += 1
                continue
            want = {"fill": fill, "color": color}
            if text is not None:
                want["text"] = text
            bad = {k: (got[k], v) for k, v in want.items() if got[k] != v}
            if bad:
                print(f"FAIL {name} {sheet}!R{r}C{c}: " +
                      ", ".join(f"{k} got {g!r} want {w!r}" for k, (g, w) in bad.items()))
                failures += 1
        warns = doc.get("warnings", [])
        for pat in EXPECTED_WARNINGS[name]:
            if not any(re.search(pat, w, re.I) for w in warns):
                print(f"FAIL {name}: no warning matching /{pat}/ in {warns}")
                failures += 1
        if not EXPECTED_WARNINGS[name] and warns:
            print(f"FAIL {name}: unexpected warnings {warns}")
            failures += 1
        print(f"{'ok  ' if failures == 0 else '....'} {name}: {len(exp)} cells checked")
    print(f"\n{failures} failure(s)")
    return 1 if failures else 0


EXPECTATIONS: dict[str, list] = {}


def main() -> int:
    OUT.mkdir(parents=True, exist_ok=True)
    if "--verify" in sys.argv:
        # Expectations are derived from building to a scratch path, so the
        # theme they read is openpyxl's — the same one the fixture carries.
        import tempfile
        with tempfile.TemporaryDirectory() as d:
            for name, build in FIXTURES.items():
                EXPECTATIONS[name] = build(Path(d) / name)
        return verify()
    for name, build in FIXTURES.items():
        exp = build(OUT / name)
        print(f"wrote {OUT / name} ({len(exp)} expectations)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
