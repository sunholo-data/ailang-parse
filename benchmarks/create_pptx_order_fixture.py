# /// script
# requires-python = ">=3.10"
# dependencies = ["python-pptx>=1.0"]
# ///
"""Build the PPTX slide-order / speaker-notes fixture, and check the parser on it.

A partner's 21-slide deck parsed in the order 1,10,11,...,19,2,20,21,3,...,9
(the slide parts were never ordered, and the ZIP listing happened to be
lexicographic), and its speaker notes were appended after every slide with no
link back to the slide they belong to. No fixture caught either: every pptx in
data/test_files had fewer than ten slides, stored in numeric order, with
notesSlideN always belonging to slideN.

This fixture breaks each of those coincidences on purpose:

  * 13 slides, so a lexicographic order is visibly wrong (slide10 < slide2);
  * the ZIP stores its entries shuffled, so the listing order means nothing;
  * ppt/presentation.xml's sldIdLst orders the slides differently from their
    file numbers (deck position 1 is slide7.xml) — as after a user drags
    slides around in PowerPoint, which keeps the file names;
  * notes parts are created out of order (NOTES_ORDER), so notesSlideN is
    never slideN.xml's notes and never deck position N's either;
  * one slide has no notes, one slide is empty but HAS notes, one notes page
    keeps its text in a plain text box (no body placeholder), one inside a
    group shape, one spans two paragraphs.

Expectations are computed here, from DECK and FILE_NUMBER below, never from
parser output: a golden generated from our own output would only prove the
parser agrees with itself.

Usage:
    uv run benchmarks/create_pptx_order_fixture.py            # (re)build fixture
    uv run benchmarks/create_pptx_order_fixture.py --verify   # parse + compare
"""

from __future__ import annotations

import io
import json
import os
import random
import re
import subprocess
import sys
import zipfile
from pathlib import Path

from pptx import Presentation

REPO = Path(__file__).resolve().parent.parent
FIXTURE = REPO / "data" / "test_files" / "pptx_slide_order_notes.pptx"

# Deck in PRESENTATION order: (title, body, notes, notes_style).
# title None = an empty slide. notes None = no notes page at all.
DECK = [
    ("Kickoff", "Why we are here", "Open with the customer quote", "body"),
    ("Agenda", "Five topics", "Keep this under a minute", "body"),
    ("Market", "Three segments", "Segment two is the growth story", "textbox"),
    ("Pricing", "Per-page tiers", None, None),
    ("Roadmap", "Q1 to Q4", "Line one of the roadmap notes\nLine two of the roadmap notes", "body"),
    (None, None, "This slide is intentionally blank; pause here", "body"),
    ("Team", "Who does what", "Introduce the new hires", "group"),
    ("Risks", "What could go wrong", "Be candid about hiring risk", "body"),
    ("Budget", "Spend by quarter", "Budget is draft until board sign-off", "body"),
    ("Metrics", "What we measure", "Tie each metric to a decision", "body"),
    ("Timeline", "Milestones", "Milestone three moved a week", "body"),
    ("Asks", "What we need", "Make the ask explicit", "body"),
    ("Close", "Thank you", "Take questions for ten minutes", "body"),
]

# File number of each deck position's slide part (slideN.xml). Deck position 1
# lives in slide7.xml, and so on: a permutation that is neither numeric nor
# lexicographic order.
FILE_NUMBER = [7, 2, 11, 1, 13, 4, 9, 3, 12, 6, 10, 5, 8]
assert sorted(FILE_NUMBER) == list(range(1, len(DECK) + 1))

# Deck positions (1-based) in the order their notes pages are created, so the
# Kth entry's notes land in notesSlideK.xml. Chosen so that K is never the
# slide's deck position nor its file number.
NOTES_ORDER = [12, 3, 13, 9, 2, 5, 8, 1, 10, 7, 6, 11]
assert sorted(NOTES_ORDER) == [p for p, d in enumerate(DECK, 1) if d[2] is not None]


def build(path: Path) -> None:
    prs = Presentation()
    prs.slides  # python-pptx renames slide parts on first access; do it while empty
    title_body = prs.slide_layouts[1]
    blank = prs.slide_layouts[6]

    # Create slides in FILE order, so slideN.xml holds the position whose
    # file number is N.
    by_file = {n: pos for pos, n in enumerate(FILE_NUMBER)}
    slides_by_pos: dict[int, object] = {}
    for n in range(1, len(DECK) + 1):
        pos = by_file[n]
        title, body, _, _ = DECK[pos]
        if title is None:
            slide = prs.slides.add_slide(blank)
        else:
            slide = prs.slides.add_slide(title_body)
            slide.shapes.title.text = title
            slide.placeholders[1].text = body
        slides_by_pos[pos] = slide

    # Notes pages in NOTES_ORDER: notesSlide1 is deck position 12's.
    for p in NOTES_ORDER:
        slides_by_pos[p - 1].notes_slide.notes_text_frame.text = DECK[p - 1][2]

    # sldIdLst into deck order. The elements currently sit in file order.
    lst = prs.slides._sldIdLst
    ids = list(lst)
    for el in ids:
        lst.remove(el)
    for pos in range(len(DECK)):
        lst.append(ids[FILE_NUMBER[pos] - 1])

    buf = io.BytesIO()
    prs.save(buf)
    parts = {i.filename: zipfile.ZipFile(buf).read(i.filename)
             for i in zipfile.ZipFile(buf).infolist()}

    # Notes stored outside the body placeholder, the two ways writers do it.
    notes_part = notes_part_for(parts)
    for pos, (_, _, notes, style) in enumerate(DECK):
        if style == "textbox":
            name = notes_part[pos]
            parts[name] = re.sub(rb'<p:ph type="body"[^>]*/>', b"", parts[name], count=1)
        elif style == "group":
            name = notes_part[pos]
            parts[name] = wrap_body_in_group(parts[name])

    # Store the entries shuffled — [Content_Types].xml stays first.
    names = [n for n in parts if n != "[Content_Types].xml"]
    random.Random(21).shuffle(names)
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for n in ["[Content_Types].xml"] + names:
            info = zipfile.ZipInfo(n, date_time=(2026, 9, 29, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            z.writestr(info, parts[n])


def notes_part_for(parts: dict[str, bytes]) -> dict[int, str]:
    """Deck position -> notes part name, resolved through the slide rels."""
    out = {}
    for pos, n in enumerate(FILE_NUMBER):
        rels = parts[f"ppt/slides/_rels/slide{n}.xml.rels"].decode()
        m = re.search(r'Type="[^"]*/notesSlide"[^>]*Target="\.\./([^"]+)"', rels) or \
            re.search(r'Target="\.\./([^"]+)"[^>]*Type="[^"]*/notesSlide"', rels)
        if m:
            out[pos] = "ppt/" + m.group(1)
    return out


def wrap_body_in_group(xml: bytes) -> bytes:
    s = xml.decode()
    sps = list(re.finditer(r"<p:sp>.*?</p:sp>", s, re.S))
    body = next(m for m in sps if 'type="body"' in m.group(0))
    group = ('<p:grpSp><p:nvGrpSpPr><p:cNvPr id="99" name="Notes Group"/>'
             '<p:cNvGrpSpPr/><p:nvPr/></p:nvGrpSpPr><p:grpSpPr/>'
             + body.group(0) + "</p:grpSp>")
    return (s[:body.start()] + group + s[body.end():]).encode()


# --- independent oracle ------------------------------------------------------

def expected_sections() -> list[tuple[str, str, str]]:
    """(kind, name, text) for every top-level section, in order."""
    out = []
    for pos, (title, body, notes, _) in enumerate(DECK, start=1):
        text = "" if title is None else f"{title}|{body}"
        out.append(("slide", "", text))
        if notes is not None:
            out.append(("notes", f"Slide {pos}", notes))
    return out


def actual_sections(doc: dict) -> list[tuple[str, str, str]]:
    out = []
    for b in doc["document"]["blocks"]:
        if b.get("type") != "section":
            continue
        texts = [c.get("text", "") for c in b.get("blocks", [])]
        out.append((b.get("kind", ""), b.get("name", ""), "|".join(texts)))
    return out


def check_structure() -> int:
    """The fixture really has the properties it claims."""
    z = zipfile.ZipFile(FIXTURE)
    names = [n for n in z.namelist() if re.fullmatch(r"ppt/slides/slide\d+\.xml", n)]
    nums = [int(re.search(r"(\d+)", n.rsplit("/", 1)[1]).group(1)) for n in names]
    failures = 0
    if nums == sorted(nums):
        print("FAIL fixture: slide parts are stored in numeric order")
        failures += 1
    parts = {n: z.read(n) for n in z.namelist()}
    for pos, name in notes_part_for(parts).items():
        if re.search(r"(\d+)", name.rsplit("/", 1)[1]).group(1) in (str(pos + 1), str(FILE_NUMBER[pos])):
            print(f"FAIL fixture: {name} numbering matches its slide")
            failures += 1
    return failures


def verify() -> int:
    outdir = REPO / "docparse" / "data"
    subprocess.run(
        ["ailang", "run", "--entry", "main", "--caps", "IO,FS,Env",
         "--max-recursion-depth", "50000", "docparse/main.ail", str(FIXTURE)],
        cwd=REPO, check=True, capture_output=True,
        env={**os.environ, "DOCPARSE_OUTPUT_DIR": str(outdir)})
    failures = check_structure()
    doc = json.load(open(outdir / f"{FIXTURE.name}.json"))
    want, got = expected_sections(), actual_sections(doc)
    if want != got:
        failures += 1
        print(f"FAIL sections differ ({len(got)} got, {len(want)} want):")
        for i in range(max(len(want), len(got))):
            w = want[i] if i < len(want) else None
            g = got[i] if i < len(got) else None
            print(f"  {'  ' if w == g else '!!'} want {w}\n     got  {g}")
    md = (outdir / f"{FIXTURE.name}.md").read_text()
    for pos, (title, _, notes, _) in enumerate(DECK, start=1):
        if notes is None:
            continue
        label = f"*Speaker notes (Slide {pos}):*"
        at = md.find(label)
        if at < 0:
            print(f"FAIL markdown: no '{label}'")
            failures += 1
        elif title is not None and md.rfind(f"# {title}", 0, at) < 0:
            print(f"FAIL markdown: '{label}' does not follow '# {title}'")
            failures += 1
    print(f"{'ok  ' if failures == 0 else '....'} {FIXTURE.name}: "
          f"{len(want)} sections checked")
    print(f"\n{failures} failure(s)")
    return 1 if failures else 0


def main() -> int:
    if "--verify" in sys.argv:
        return verify()
    build(FIXTURE)
    print(f"wrote {FIXTURE} ({len(expected_sections())} expected sections)")
    return check_structure()


if __name__ == "__main__":
    sys.exit(main())
